"""Parsing attendee spreadsheets (.csv / .xlsx) into validated user records.

Real spreadsheets are messy: headers are named whatever the organiser felt like,
Excel turns phone numbers into floats, and Windows CSV exports carry a BOM.
Everything defensive about that lives here so the routes stay readable.
"""

import csv
import io
import re
from typing import Iterable

from openpyxl import load_workbook
from pydantic import ValidationError

from app.models import ImportRow, UserCreate

MAX_ROWS = 5000

# Header aliases, normalised to lowercase alphanumerics before lookup.
# Covers the spellings that actually turn up in event spreadsheets.
COLUMN_ALIASES: dict[str, str] = {
    "name": "name",
    "fullname": "name",
    "attendeename": "name",
    "participantname": "name",
    "firstname": "name",
    "email": "email",
    "emailaddress": "email",
    "emailid": "email",
    "mail": "email",
    "phone": "phone",
    "phonenumber": "phone",
    "mobile": "phone",
    "mobilenumber": "phone",
    "contact": "phone",
    "contactnumber": "phone",
    "organization": "organization",
    "organisation": "organization",
    "company": "organization",
    "college": "organization",
    "institution": "organization",
    "org": "organization",
    "affiliation": "organization",
}


def _normalize_header(header: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(header or "").strip().lower())


def _fuzzy_field(normalized: str) -> str | None:
    """Fallback for headers the alias table does not list verbatim.

    Enumerating every spelling an organiser might use ("Mobile No", "Contact #",
    "Cell") is a losing game, so unrecognised headers fall through to substring
    rules. Order matters: "Company Name" must resolve to organization, not name,
    so the more specific patterns are checked first.
    """
    for needle, field in (
        ("mail", "email"),
        ("compan", "organization"),
        ("organi", "organization"),
        ("college", "organization"),
        ("institut", "organization"),
        ("employer", "organization"),
        ("firm", "organization"),
        ("phone", "phone"),
        ("mobile", "phone"),
        ("contact", "phone"),
        ("cell", "phone"),
        ("whatsapp", "phone"),
        ("name", "name"),
    ):
        if needle in normalized:
            return field
    return None


def map_columns(headers: Iterable[str]) -> tuple[dict[int, str], list[str]]:
    """Map column positions to our field names.

    Exact aliases win; anything left over gets the substring fallback. First
    match claims a field so a sheet with both "Name" and "Full Name" does not
    fight with itself.
    """
    normalized = [_normalize_header(h) for h in headers]
    mapping: dict[int, str] = {}
    claimed: set[str] = set()

    # Pass 1: exact alias matches, which are unambiguous.
    for idx, norm in enumerate(normalized):
        field = COLUMN_ALIASES.get(norm)
        if field and field not in claimed:
            mapping[idx] = field
            claimed.add(field)

    # Pass 2: substring guesses for whatever is still unclaimed.
    for idx, norm in enumerate(normalized):
        if idx in mapping or not norm:
            continue
        field = _fuzzy_field(norm)
        if field and field not in claimed:
            mapping[idx] = field
            claimed.add(field)

    unmapped = [
        str(h).strip()
        for idx, h in enumerate(headers)
        if idx not in mapping and str(h or "").strip()
    ]
    return mapping, unmapped


def _clean_cell(value: object) -> str:
    """Coerce a spreadsheet cell to a clean string.

    Excel stores a phone number as a float, so 9840011223 comes back as
    9840011223.0, and long digit strings arrive in scientific notation. Both
    would end up embedded in a QR code, so they are fixed here.
    """
    if value is None:
        return ""
    if isinstance(value, float):
        # Whole-number floats are almost always integers Excel widened.
        if value.is_integer():
            return str(int(value))
        return str(value)
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _decode_csv(raw: bytes) -> str:
    """Decode CSV bytes, tolerating Excel's Windows exports.

    utf-8-sig strips the BOM Excel writes; cp1252 is the usual fallback for
    files saved on Windows with accented names in them.
    """
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _read_csv(raw: bytes) -> list[list[str]]:
    text = _decode_csv(raw)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel  # single-column files defeat the sniffer
    return [[_clean_cell(c) for c in row] for row in csv.reader(io.StringIO(text), dialect)]


def _read_xlsx(raw: bytes) -> list[list[str]]:
    # read_only keeps memory flat on large sheets; data_only takes the cached
    # value of a formula rather than the formula text.
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        return [[_clean_cell(c) for c in row] for row in wb.active.iter_rows(values_only=True)]
    finally:
        wb.close()


def read_table(filename: str, raw: bytes) -> list[list[str]]:
    """Dispatch on file extension and return raw rows including the header."""
    lower = filename.lower()
    if lower.endswith(".csv"):
        rows = _read_csv(raw)
    elif lower.endswith(".xlsx"):
        rows = _read_xlsx(raw)
    elif lower.endswith(".xls"):
        raise ValueError(
            "Legacy .xls files are not supported. Open the file and re-save it as .xlsx or .csv."
        )
    else:
        raise ValueError(f"Unsupported file type '{filename}'. Upload a .csv or .xlsx file.")

    # Excel routinely reports thousands of trailing rows that are entirely blank.
    while rows and not any(cell for cell in rows[-1]):
        rows.pop()

    if not rows:
        raise ValueError("The file appears to be empty.")
    return rows


def parse_rows(rows: list[list[str]], existing_emails: set[str]) -> tuple[list[ImportRow], dict[int, str], list[str]]:
    """Validate each data row and classify what should happen to it."""
    header, *data_rows = rows
    mapping, unmapped = map_columns(header)

    if "name" not in mapping.values() or "email" not in mapping.values():
        found = ", ".join(h for h in header if h) or "none"
        raise ValueError(
            "Could not find a name column and an email column. "
            f"Columns found: {found}. Rename them to 'Name' and 'Email' and try again."
        )

    if len(data_rows) > MAX_ROWS:
        raise ValueError(f"File has {len(data_rows)} rows; the limit is {MAX_ROWS}.")

    results: list[ImportRow] = []
    seen_in_file: set[str] = set()

    for offset, row in enumerate(data_rows):
        row_number = offset + 2  # 1-indexed, and the header occupies row 1
        raw_values = {
            field: (row[idx] if idx < len(row) else "") for idx, field in mapping.items()
        }

        if not any(raw_values.values()):
            continue  # blank row in the middle of the sheet

        try:
            user = UserCreate(
                name=raw_values.get("name", ""),
                email=raw_values.get("email", ""),
                phone=raw_values.get("phone") or None,
                organization=raw_values.get("organization") or None,
            )
        except ValidationError as e:
            first = e.errors()[0]
            field = first["loc"][0] if first["loc"] else "row"
            results.append(
                ImportRow(
                    row_number=row_number,
                    status="invalid",
                    raw=raw_values,
                    error=f"{field}: {first['msg']}",
                )
            )
            continue

        if user.email in seen_in_file:
            status = "duplicate_in_file"
        elif user.email in existing_emails:
            status = "already_exists"
        else:
            status = "ok"
            seen_in_file.add(user.email)

        results.append(
            ImportRow(row_number=row_number, status=status, data=user, raw=raw_values)
        )

    return results, mapping, unmapped
